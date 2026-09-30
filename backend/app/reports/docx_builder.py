import os
from pathlib import Path
from xml.sax.saxutils import escape
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls
from docx.opc.constants import RELATIONSHIP_TYPE



def safe_url(url: str) -> str:
    if not url or not isinstance(url, str):
        return ""
    url = url.strip()
    if url.startswith(("http://", "https://", "mailto:")):
        return url
    return f"https://{url}"


def add_hyperlink(paragraph, url: str, text: str, color="0284C7", underline=True):
    url = safe_url(url)
    if not url:
        paragraph.add_run(text)
        return
    try:
        part = paragraph.part
        r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = parse_xml(
            f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'r:id="{r_id}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"/>'
        )
        new_run = parse_xml(f'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        u_val = "single" if underline else "none"
        rPr = parse_xml(
            f'<w:rPr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:color w:val="{color}"/><w:u w:val="{u_val}"/></w:rPr>'
        )
        new_run.append(rPr)
        new_run_text = parse_xml(
            f'<w:t xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">{escape(text)}</w:t>'
        )
        new_run.append(new_run_text)
        hyperlink.append(new_run)
        paragraph._p.append(hyperlink)
    except Exception:
        run = paragraph.add_run(text)
        run.font.color.rgb = RGBColor(2, 132, 199)
        run.font.underline = True


def set_cell_background(cell, fill_hex: str):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def generate_docx_report(db, run_id: str, output_path) -> str:
    from app.reports.builder import load_report
    data = load_report(db, run_id)
    doc = docx.Document()

    # Page Margins (0.8 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Base style font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    # Header / Eyebrow
    p_eyebrow = doc.add_paragraph()
    p_eyebrow.paragraph_format.space_before = Pt(0)
    p_eyebrow.paragraph_format.space_after = Pt(2)
    run_eye = p_eyebrow.add_run("COMPETITIVE INTELLIGENCE EXECUTIVE REPORT")
    run_eye.bold = True
    run_eye.font.size = Pt(8.5)
    run_eye.font.color.rgb = RGBColor(2, 132, 199)

    # Company Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(2)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(data["company"]["company_name"])
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    # Industry / Business Model
    ind = data["company"].get("industry", "Not verified")
    model = data["company"].get("business_model", "Not verified")
    sub_parts = []
    if ind and ind != "Not verified":
        sub_parts.append(ind)
    if model and model != "Not verified":
        sub_parts.append(model)
    if sub_parts:
        p_sub = doc.add_paragraph(" | ".join(sub_parts))
        p_sub.paragraph_format.space_after = Pt(4)
        p_sub.runs[0].font.size = Pt(11)
        p_sub.runs[0].font.color.rgb = RGBColor(71, 85, 105)

    # Website Link
    site_url = data["company"].get("website")
    if site_url:
        p_site = doc.add_paragraph()
        p_site.paragraph_format.space_after = Pt(14)
        add_hyperlink(p_site, site_url, site_url)

    # Executive Brief Box
    p_brief_hdr = doc.add_paragraph()
    p_brief_hdr.paragraph_format.space_before = Pt(8)
    p_brief_hdr.paragraph_format.space_after = Pt(2)
    run_brief_hdr = p_brief_hdr.add_run("Executive Brief")
    run_brief_hdr.bold = True
    run_brief_hdr.font.size = Pt(12)
    run_brief_hdr.font.color.rgb = RGBColor(15, 23, 42)

    desc_text = data["company"].get("description") or "Verified competitive landscape, product catalog extraction, and market positioning analysis."
    p_desc = doc.add_paragraph(desc_text)
    p_desc.paragraph_format.space_after = Pt(8)

    # Summary Stats
    p_stats = doc.add_paragraph()
    p_stats.paragraph_format.space_after = Pt(16)
    stats_list = [f"Crawl Sample: {len(data.get('pages', []))} pages"]
    if data.get("products"):
        stats_list.append(f"Discovered Products: {len(data['products'])}")
    if data.get("competitors"):
        stats_list.append(f"Verified Competitors: {len(data['competitors'])}")
    if data.get("keywords"):
        stats_list.append(f"Search Hypotheses: {len(data['keywords'])}")
    run_stats = p_stats.add_run(" • ".join(stats_list))
    run_stats.font.size = Pt(9)
    run_stats.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 01 / CATALOG & OFFERINGS
    if data.get("products") or data.get("categories"):
        h1 = doc.add_heading("01 / Discovered Catalog & Offerings", level=2)
        h1.paragraph_format.space_before = Pt(14)
        h1.paragraph_format.space_after = Pt(6)

        price_summary = data.get("price_range_summary")
        if price_summary and price_summary != "Not publicly disclosed":
            p_price = doc.add_paragraph()
            p_price.paragraph_format.space_after = Pt(6)
            r1 = p_price.add_run("Price Range Analysis: ")
            r1.bold = True
            p_price.add_run(price_summary)

        if data.get("products"):
            doc.add_paragraph("Sampled Products:").runs[0].bold = True
            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False

            col_widths = [Inches(3.6), Inches(1.6), Inches(1.6)]
            hdr_cells = table.rows[0].cells
            headers = ["Product Name", "Price / Currency", "Availability"]
            for i, h in enumerate(headers):
                hdr_cells[i].text = h
                hdr_cells[i].width = col_widths[i]
                set_cell_background(hdr_cells[i], "0F172A")
                set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
                for p in hdr_cells[i].paragraphs:
                    p.runs[0].font.bold = True
                    p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                    p.runs[0].font.size = Pt(9)

            for prod in data["products"]:
                row_cells = table.add_row().cells
                for i, w in enumerate(col_widths):
                    row_cells[i].width = w
                    set_cell_margins(row_cells[i], top=90, bottom=90, left=150, right=150)

                row_cells[0].paragraphs[0].add_run(prod.get("product_name", "Unknown")).bold = True
                row_cells[0].paragraphs[0].runs[0].font.size = Pt(9)

                price_str = f"{prod.get('price', 'Not disclosed')} {prod.get('currency', '')}".strip()
                row_cells[1].paragraphs[0].add_run(price_str).font.size = Pt(9)

                avail = (prod.get("availability") or "Available").replace("https://schema.org/", "")
                row_cells[2].paragraphs[0].add_run(avail).font.size = Pt(9)

            doc.add_paragraph().paragraph_format.space_after = Pt(8)

        if data.get("categories"):
            p_cat = doc.add_paragraph()
            p_cat.paragraph_format.space_before = Pt(6)
            p_cat.paragraph_format.space_after = Pt(12)
            p_cat.add_run("Categories & Collections: ").bold = True
            cat_links = []
            for c in data["categories"]:
                cat_links.append(c.get("category_name", "Category"))
            p_cat.add_run(", ".join(cat_links))

    # 02 / VERIFIED COMPETITORS
    if data.get("competitors"):
        h2 = doc.add_heading("02 / Verified Competitors", level=2)
        h2.paragraph_format.space_before = Pt(16)
        h2.paragraph_format.space_after = Pt(4)

        p_c_intro = doc.add_paragraph("Competitors independently crawled and qualified against offering and industry overlap:")
        p_c_intro.paragraph_format.space_after = Pt(8)
        p_c_intro.runs[0].font.color.rgb = RGBColor(100, 116, 139)

        for comp in data["competitors"]:
            p_comp = doc.add_paragraph()
            p_comp.paragraph_format.space_before = Pt(6)
            p_comp.paragraph_format.space_after = Pt(2)
            r_name = p_comp.add_run(comp.get("competitor_name", "Competitor"))
            r_name.bold = True
            r_name.font.size = Pt(11)
            r_type = p_comp.add_run(f"  [{comp.get('competitor_type', 'DIRECT')}]")
            r_type.font.size = Pt(8.5)
            r_type.font.color.rgb = RGBColor(2, 132, 199)

            comp_site = comp.get("website") or comp.get("domain") or ""
            if comp_site:
                p_link = doc.add_paragraph()
                p_link.paragraph_format.space_after = Pt(2)
                add_hyperlink(p_link, comp_site, comp.get("domain") or comp_site)

            desc = comp.get("description")
            if desc and desc not in ("Not found", "Not verified", ""):
                p_cdesc = doc.add_paragraph(desc)
                p_cdesc.paragraph_format.space_after = Pt(2)
                p_cdesc.runs[0].font.size = Pt(9)
                p_cdesc.runs[0].font.color.rgb = RGBColor(71, 85, 105)

            # Details
            prods = comp.get("products", [])
            if prods:
                p_prods = doc.add_paragraph()
                p_prods.paragraph_format.space_after = Pt(1)
                p_prods.add_run("Discovered Offerings: ").bold = True
                p_prods.add_run(", ".join(map(str, prods[:6])))
                p_prods.runs[0].font.size = Pt(8.5)
                p_prods.runs[1].font.size = Pt(8.5)

            pricing = comp.get("pricing_model")
            if pricing and pricing != "Not publicly disclosed":
                p_price = doc.add_paragraph()
                p_price.paragraph_format.space_after = Pt(4)
                p_price.add_run("Pricing: ").bold = True
                p_price.add_run(pricing)
                p_price.runs[0].font.size = Pt(8.5)
                p_price.runs[1].font.size = Pt(8.5)

    # 03 / KEYWORDS & SEARCH INTENT
    if data.get("keywords"):
        h3 = doc.add_heading("03 / Keywords & Search Intent", level=2)
        h3.paragraph_format.space_before = Pt(16)
        h3.paragraph_format.space_after = Pt(6)

        table_kw = doc.add_table(rows=1, cols=3)
        table_kw.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_kw.autofit = False

        kw_widths = [Inches(3.8), Inches(1.8), Inches(1.2)]
        kw_hdr = table_kw.rows[0].cells
        for i, h in enumerate(["Search Hypothesis", "Inferred Intent", "Intent Type"]):
            kw_hdr[i].text = h
            kw_hdr[i].width = kw_widths[i]
            set_cell_background(kw_hdr[i], "0F172A")
            set_cell_margins(kw_hdr[i], top=120, bottom=120, left=150, right=150)
            for p in kw_hdr[i].paragraphs:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                p.runs[0].font.size = Pt(9)

        for kw in data["keywords"][:15]:
            row_cells = table_kw.add_row().cells
            for i, w in enumerate(kw_widths):
                row_cells[i].width = w
                set_cell_margins(row_cells[i], top=80, bottom=80, left=150, right=150)

            row_cells[0].paragraphs[0].add_run(kw.get("keyword", "")).bold = True
            row_cells[0].paragraphs[0].runs[0].font.size = Pt(9)

            row_cells[1].paragraphs[0].add_run(kw.get("intent", "")).font.size = Pt(9)

            itype = kw.get("intent_type") or "Product"
            row_cells[2].paragraphs[0].add_run(itype).font.size = Pt(9)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 04 / SOCIAL MEDIA CHANNELS (Only if verified on site)
    if data.get("social_profiles") and len(data["social_profiles"]) > 0:
        h4 = doc.add_heading("04 / Official Social Media Channels", level=2)
        h4.paragraph_format.space_before = Pt(16)
        h4.paragraph_format.space_after = Pt(6)

        for sp in data["social_profiles"]:
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_after = Pt(3)
            p_sp.add_run(f"• {sp.get('platform')}: ").bold = True
            p_sp.runs[0].font.size = Pt(9)
            add_hyperlink(p_sp, sp.get("profile_url", ""), sp.get("profile_url", ""))

    # 05 / VERIFIED EVIDENCE & SOURCES
    if data.get("sources"):
        h5 = doc.add_heading("05 / Verified Evidence & Sources", level=2)
        h5.paragraph_format.space_before = Pt(16)
        h5.paragraph_format.space_after = Pt(6)

        for idx, src in enumerate(data["sources"][:15], 1):
            p_src = doc.add_paragraph()
            p_src.paragraph_format.space_after = Pt(3)
            r_num = p_src.add_run(f"[{idx}] ")
            r_num.bold = True
            r_num.font.color.rgb = RGBColor(2, 132, 199)
            r_num.font.size = Pt(8.5)
            r_title = p_src.add_run(f"{src.get('page_title', 'Source Page')} — ")
            r_title.bold = True
            r_title.font.size = Pt(8.5)
            add_hyperlink(p_src, src.get("url", ""), src.get("url", ""))

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return str(output_path)
